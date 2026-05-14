from __future__ import annotations

import csv
import json
import re
import sys
import zipfile
from collections import Counter, defaultdict
from pathlib import Path
from xml.etree import ElementTree as ET

from PIL import Image, ImageOps


# The Codex document runtime supplies python-docx. Keep this path explicit so
# report generation does not depend on the project virtualenv package set.
DOCX_SITE_PACKAGES = Path(
    "/Users/bhuwanjungthapa/.cache/codex-runtimes/"
    "codex-primary-runtime/dependencies/python/lib/python3.12/site-packages"
)
if DOCX_SITE_PACKAGES.exists():
    sys.path.append(str(DOCX_SITE_PACKAGES))

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
REPORTS = ROOT / "reports"
FIGURES = REPORTS / "figures"
ASSETS = REPORTS / "final_report_assets"
OUT = REPORTS / "Final_Project_Report.docx"

BLACK = RGBColor(0, 0, 0)
BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(11, 37, 69)
MUTED = RGBColor(88, 96, 105)
LIGHT_GRAY = "F2F4F7"
BLUE_GRAY = "E8EEF5"
BORDER = "C8D0DA"


def read_json(path: Path) -> dict:
    with path.open() as f:
        return json.load(f)


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(newline="") as f:
        return list(csv.DictReader(f))


def extract_pptx_slide_text(path: Path) -> dict[int, list[str]]:
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}

    def num_key(name: str) -> int:
        match = re.search(r"slide(\d+)\.xml$", name)
        return int(match.group(1)) if match else 999

    slides: dict[int, list[str]] = {}
    with zipfile.ZipFile(path) as zf:
        names = sorted(
            [n for n in zf.namelist() if re.match(r"ppt/slides/slide\d+\.xml$", n)],
            key=num_key,
        )
        for name in names:
            root = ET.fromstring(zf.read(name))
            lines: list[str] = []
            for paragraph in root.findall(".//a:p", ns):
                text = "".join(t.text or "" for t in paragraph.findall(".//a:t", ns)).strip()
                if text:
                    lines.append(text)
            slides[num_key(name)] = lines
    return slides


def load_metrics() -> dict:
    summary = read_json(REPORTS / "evaluation_summary.json")
    results = read_csv(REPORTS / "evaluation_results.csv")
    ranks = read_csv(FIGURES / "ranks_debug.csv")
    manifest = read_json(ROOT / "indexes" / "manifest.json")
    lora_train = read_json(
        ROOT / "experiments/lora_qlora/adapters/flan-t5-small-lora/training_summary.json"
    )
    lora_eval = read_json(
        ROOT / "experiments/lora_qlora/results/adapter_eval_summary.json"
    )
    labels = Counter(row["label"] for row in results)
    grounded = Counter(row["grounded_or_not"] for row in results)

    hit_sweep = {}
    n_ranks = len(ranks)
    for k in (1, 3, 5, 10):
        hits = 0
        for row in ranks:
            rank = row.get("gold_rank", "")
            if rank and float(rank) <= k:
                hits += 1
        hit_sweep[k] = {"hits": hits, "n": n_ranks, "rate": hits / n_ranks if n_ranks else 0.0}

    heat = defaultdict(Counter)
    for idx, row in enumerate(results):
        rank = ranks[idx].get("gold_rank", "") if idx < len(ranks) else ""
        hit5 = bool(rank and float(rank) <= 5)
        heat["Hit@5 = True" if hit5 else "Hit@5 = False"][row["label"]] += 1

    return {
        "summary": summary,
        "results": results,
        "ranks": ranks,
        "manifest": manifest,
        "lora_train": lora_train,
        "lora_eval": lora_eval,
        "labels": labels,
        "grounded": grounded,
        "hit_sweep": hit_sweep,
        "heat": heat,
        "pptx_slides": extract_pptx_slide_text(REPORTS / "Final_Project_Presentation.pptx"),
    }


def load_runtime_notes() -> dict[str, str]:
    notes = {
        "gemini_model": "Gemini Lite preview model configured in .env",
        "gemini_key": "not checked",
        "openai_key": "not checked",
    }
    env_path = ROOT / ".env"
    if not env_path.exists():
        return notes
    for line in env_path.read_text().splitlines():
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        value = value.strip().strip("\"'")
        if key == "GEMINI_MODEL" and value:
            notes["gemini_model"] = value
        elif key == "GEMINI_API_KEY":
            notes["gemini_key"] = "set" if value else "empty"
        elif key == "OPENAI_API_KEY":
            notes["openai_key"] = "set" if value else "empty"
    return notes


def pct(value: float, digits: int = 0) -> str:
    return f"{value * 100:.{digits}f}%"


def fmt(value: float, digits: int = 3) -> str:
    return f"{value:.{digits}f}"


def build_composite(
    image_names: list[str],
    output_name: str,
    labels: list[str] | None = None,
    gap: int = 42,
    pad: int = 26,
) -> Path:
    ASSETS.mkdir(parents=True, exist_ok=True)
    imgs = [Image.open(FIGURES / name).convert("RGB") for name in image_names]
    target_h = min(img.height for img in imgs)
    resized = []
    for img in imgs:
        w = int(img.width * (target_h / img.height))
        resized.append(img.resize((w, target_h), Image.LANCZOS))

    label_h = 34 if labels else 0
    width = sum(img.width for img in resized) + gap * (len(resized) - 1) + pad * 2
    height = target_h + label_h + pad * 2
    canvas = Image.new("RGB", (width, height), "white")
    x = pad
    for idx, img in enumerate(resized):
        if labels:
            # Simple text-free label block: the Word caption names each panel.
            pass
        canvas.paste(ImageOps.expand(img, border=1, fill=(220, 224, 230)), (x, pad + label_h))
        x += img.width + gap

    out = ASSETS / output_name
    canvas.save(out, quality=95)
    return out


def setup_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    title = styles["Title"]
    title.font.name = "Calibri"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    title.font.size = Pt(22)
    title.font.bold = True
    title.font.color.rgb = INK
    title.paragraph_format.space_after = Pt(4)

    subtitle = styles["Subtitle"]
    subtitle.font.name = "Calibri"
    subtitle._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    subtitle._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    subtitle.font.size = Pt(13)
    subtitle.font.color.rgb = MUTED
    subtitle.paragraph_format.space_after = Pt(12)

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for list_style in ["List Bullet", "List Number"]:
        style = styles[list_style]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.167


def set_run_font(run, size: float | None = None, color: RGBColor | None = None, bold=None, italic=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_run_font(run, size=9, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(t)
    fld.append(r)
    paragraph._p.append(fld)


def set_header_footer(doc: Document) -> None:
    section = doc.sections[0]
    header = section.header.paragraphs[0]
    header.text = "Generative AI RAG Document Q&A Chatbot"
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in header.runs:
        set_run_font(run, size=9, color=MUTED)
    add_page_number(section.footer.paragraphs[0])


def paragraph_bottom_border(paragraph, color: str = BORDER, size: str = "8") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def add_masthead(doc: Document, metrics: dict) -> None:
    p = doc.add_paragraph("Final Project Report", style="Title")
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.keep_with_next = True

    p = doc.add_paragraph("Generative AI RAG Document Q&A Chatbot", style="Subtitle")
    p.paragraph_format.keep_with_next = True

    rows = [
        ("Course", "CSE 434 / CSE 534 — Generative AI"),
        ("Team Member", "Bhuwan Jung Thapa"),
        ("Report Date", "May 14, 2026"),
        ("Assignment Due", "May 15, 2026"),
    ]
    for label, value in rows:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(f"{label}: ")
        set_run_font(r, bold=True)
        r = p.add_run(value)
        set_run_font(r)

    rule = doc.add_paragraph()
    rule.paragraph_format.space_after = Pt(12)
    paragraph_bottom_border(rule, "9FB6CF", "10")

    lead = doc.add_paragraph()
    lead.paragraph_format.space_after = Pt(10)
    r = lead.add_run("Project focus. ")
    set_run_font(r, bold=True, color=INK)
    lead.add_run(
        "This report is about the PDF problem I kept running into as a student: too many course files, spread across too many semesters, and no reliable way to ask questions across all of them without losing track or getting hallucinated answers."
    )


def set_cell(cell, text: str, bold: bool = False, shade: str | None = None, align=None) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    set_run_font(run, size=9.5, bold=bold)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    tc_pr = cell._tc.get_or_add_tcPr()
    if shade:
        shd = tc_pr.find(qn("w:shd"))
        if shd is None:
            shd = OxmlElement("w:shd")
            tc_pr.append(shd)
        shd.set(qn("w:fill"), shade)


def set_table_geometry(table, widths_in: list[float], indent_dxa: int = 120) -> None:
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(int(w * 1440) for w in widths_in)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_in:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(int(width * 1440)))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths_in[idx])
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(widths_in[idx] * 1440)))
            tc_w.set(qn("w:type"), "dxa")
            tc_mar = tc_pr.find(qn("w:tcMar"))
            if tc_mar is None:
                tc_mar = OxmlElement("w:tcMar")
                tc_pr.append(tc_mar)
            for side, val in [("top", 80), ("bottom", 80), ("start", 120), ("end", 120)]:
                el = tc_mar.find(qn(f"w:{side}"))
                if el is None:
                    el = OxmlElement(f"w:{side}")
                    tc_mar.append(el)
                el.set(qn("w:w"), str(val))
                el.set(qn("w:type"), "dxa")


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float]):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    for idx, header in enumerate(headers):
        set_cell(table.rows[0].cells[idx], header, bold=True, shade=LIGHT_GRAY, align=WD_ALIGN_PARAGRAPH.CENTER)
    for row_values in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row_values):
            align = WD_ALIGN_PARAGRAPH.CENTER if len(value) < 14 or value.endswith("%") else None
            set_cell(cells[idx], value, align=align)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, size=9, color=MUTED, italic=True)


def add_picture(doc: Document, path: Path, width: float, caption: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width))
    add_caption(doc, caption)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True


def add_para(doc: Document, text: str) -> None:
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(6)


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(text, style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(5)


def add_page_break(doc: Document) -> None:
    doc.add_page_break()


def make_report() -> None:
    metrics = load_metrics()
    runtime = load_runtime_notes()
    summary = metrics["summary"]
    manifest = metrics["manifest"]
    lora_train = metrics["lora_train"]
    lora_eval = metrics["lora_eval"]
    labels = metrics["labels"]
    hit = metrics["hit_sweep"]

    hit_calib = build_composite(["fig2_hit_at_k.png", "fig6_calibration.png"], "hit_calibration.png")
    heat_labels = build_composite(["fig5_heatmap.png", "fig3_label_distribution.png"], "heatmap_labels.png")

    doc = Document()
    setup_styles(doc)
    set_header_footer(doc)
    add_masthead(doc, metrics)

    add_heading(doc, "1. Introduction & Problem Definition", 1)
    add_para(
        doc,
        "The problem I wanted to solve came from my own semester experience. A student like me takes multiple courses across multiple semesters, and every course can have many PDFs: the syllabus, weekly lecture PDFs, assignment files, project requirement PDFs, exam updates, presentation guidelines, and other announcements. I have personally downloaded the same file more than once because I could not find the last downloaded copy when I needed it."
    )
    add_para(
        doc,
        "Using a normal chatbot such as ChatGPT or Gemini also became a hurdle. There are upload limits, so I cannot just put every course PDF from a semester into one session. Long sessions also make the answer quality feel less reliable: answers can become unnecessarily long, the model may mix up documents, and the chance of hallucination goes up when the conversation gets bigger. For course policies and deadlines, a wrong answer is not a small mistake."
    )
    add_para(
        doc,
        "My goal was to build a local document Q&A system where I can upload all course PDFs into one local dataset, keep them there for later, and ask questions whenever I need. The system should retrieve the source text, answer only from that source, cite the file and page, and refuse when the answer is not supported. In simple terms, I wanted a course-document assistant that stays useful even when the dataset grows and the session gets long."
    )
    add_para(
        doc,
        "The project connects to course topics through Retrieval-Augmented Generation, prompt engineering, evaluation, and fine-tuning. I first tried OpenAI API access, but usage limits became a problem. I then moved to Gemini 2.5 Lite during testing, and for the final local test the project was configured to use a Gemini Lite preview model."
    )

    add_page_break(doc)
    add_heading(doc, "2. Dataset", 1)
    add_para(
        doc,
        f"The indexed corpus contains {manifest['num_documents']} course PDFs, {manifest['num_pages']} extracted pages, and {manifest['num_chunks']} persisted chunks. The documents are user-uploaded files stored in the project data folder: CSE 434 / CSE 534 Syllabus, Final Project Presentation Guideline, Midterm Project Check Requirement, and Project Requirement."
    )
    dataset_rows = [
        ["CSE_434___CSE_534_Syllabus.pdf", "22", "Course policies, grading, staff, schedule, and assignments"],
        ["Final_Project_Presentation_Guideline.pdf", "4", "Final presentation timing, slide count, required slide topics"],
        ["Midterm_Project_Check_Requirement.pdf", "7", "Midterm report goals, rubric, graduate expectations"],
        ["Project_Requirement.pdf", "6", "Team size, timeline, project examples, graduate requirements"],
    ]
    add_table(doc, ["Document", "Chunks", "Role in corpus"], dataset_rows, [2.75, 0.75, 2.8])
    add_picture(
        doc,
        FIGURES / "fig1_chunks_per_doc.png",
        5.8,
        "Figure 1. Chunk count by source document, generated from the indexed corpus.",
    )
    add_para(
        doc,
        "Preprocessing follows the final implementation rather than a hypothetical dataset plan: pypdf extracts text page by page, whitespace is normalized, headings are used to split logical sections when possible, long sections are recursively split to 800 characters with 150 characters of overlap, and chunks below 250 characters are merged backward."
    )
    add_para(
        doc,
        f"The evaluation set is eval/sample_eval_questions.csv with {summary['n']} gold questions, each containing a question, gold answer, expected document, and optional page. The LoRA experiment generated {lora_train['train_examples'] + lora_train['eval_examples']} document-grounded QA pairs from the indexed chunks, split into {lora_train['train_examples']} training examples and {lora_train['eval_examples']} held-out examples."
    )
    add_para(
        doc,
        "Ethically, the project avoids web scraping and uses only user-uploaded course PDFs. Most processing, indexing, retrieval, and logging run locally. The main privacy constraint is that final generation may call Gemini or OpenAI with the retrieved context, so the app should be used only with documents the user is allowed to send to the selected API provider."
    )

    add_page_break(doc)
    add_heading(doc, "3. Methods / System Design", 1)
    add_para(
        doc,
        "The system has two synchronized flows. Ingestion converts PDFs into searchable evidence; query time embeds the user question, retrieves the most relevant chunks, checks confidence, and only then sends the context to the LLM. The same embedding model is used for chunks and questions, which keeps both sides in one vector space."
    )
    add_picture(
        doc,
        FIGURES / "architecture_diagram.png",
        6.4,
        "Figure 2. Final RAG architecture: ingestion writes the FAISS index; query-time retrieval reads from it.",
    )
    add_para(
        doc,
        "The embedding model is sentence-transformers/all-MiniLM-L6-v2, producing 384-dimensional L2-normalized vectors. FAISS IndexFlatIP performs exact inner-product search, which is appropriate for the small corpus and avoids approximate-nearest-neighbor recall tradeoffs. The app also supports BM25 and hybrid search; in hybrid mode, final_score = alpha * dense_score + (1 - alpha) * BM25_score, with alpha set to 0.65."
    )

    add_page_break(doc)
    add_heading(doc, "3. Methods / System Design (continued)", 1)
    add_para(
        doc,
        "The answer prompt is intentionally strict: answer only from the supplied context, cite the source as [filename p.PAGE], and return a fixed refusal sentence if the documents do not support the answer. A confidence guardrail runs before generation; if the top retrieved score is below 0.25, the LLM is not called and the system returns the refusal directly."
    )
    hyper_rows = [
        ["Chunk size / overlap / min chunk", "800 / 150 / 250 chars", "Balances context continuity with retrieval precision"],
        ["Embedding model", "all-MiniLM-L6-v2, 384d", "Small, fast, free baseline for local use"],
        ["Index", "FAISS IndexFlatIP", "Exact search is sufficient at 39 chunks"],
        ["Retrieval", "Hybrid FAISS + BM25, alpha=0.65", "Combines semantic paraphrase and exact course terms"],
        ["Top-k / min_score", "4 / 0.25", "Limits context size and refuses weak retrieval"],
        ["LLM layer", "Gemini Lite preview model or GPT-4o-mini", "Provider selected at runtime; OpenAI was limited, so Gemini Lite preview was used for final testing"],
        ["LoRA setup", "FLAN-T5-small, r=8, alpha=16, dropout=0.05", "Parameter-efficient style adaptation"],
        ["LoRA training", "3 epochs, batch=2, lr=5e-4", "Saved adapter keeps base model frozen"],
    ]
    add_table(doc, ["Component", "Final setting", "Reason"], hyper_rows, [1.55, 2.05, 2.7])
    add_para(
        doc,
        "Evaluation uses two passes. The full pipeline pass records predicted answer, retrieved documents, page ranges, top score, Hit@k, grounded_or_not, token overlap, semantic similarity, BERTScore, and label. A retrieval-only pass checks ranks up to k=10 so the report can separate retrieval ranking from answer generation."
    )
    add_para(
        doc,
        "For the graduate-level part, I added a LoRA experiment to see whether a small model could learn the answer style I wanted: use the given context, stay grounded, and include citations. I still kept RAG as the main system because the documents can change any time when a user uploads new PDFs. I did not run full QLoRA on my local Mac, but I plan to continue this part in Google Colab using a T4 GPU, where 4-bit QLoRA training is more realistic."
    )

    add_page_break(doc)
    add_heading(doc, "4. Final Results", 1)
    metric_rows = [
        ["Evaluation questions", str(summary["n"]), "All rows labeled"],
        ["Overall retrieval source hit", pct(summary["retrieval_hit_rate"]), "24 of 25 saved evaluation rows matched the expected source"],
        ["Correctness rate", pct(summary["correctness_rate"]), f"{labels['Correct']} Correct / {summary['n_labeled']} labeled"],
        ["Partially correct", str(labels["Partially Correct"]), "Largest remaining error type"],
        ["Unsupported refusals", str(labels["Unsupported"]), "Appropriate silence when support is missing"],
        ["Hallucination rate", pct(summary["hallucination_rate"]), "0 hallucinated labels"],
        ["Grounded-by-guardrail rate", pct(summary["grounded_rate"]), f"{metrics['grounded']['True']} grounded / {summary['n']} questions"],
        ["Semantic similarity mean", fmt(summary["semantic_similarity_mean"]), "Predicted vs. gold answer"],
        ["BERTScore F1 mean", fmt(summary["bertscore_f1_mean"]), "Mean precision 0.711, recall 0.838"],
        ["Calibration ECE", fmt(summary["calibration_ece"]), "Lower is better"],
    ]
    add_table(doc, ["Metric", "Value", "Interpretation"], metric_rows, [2.05, 1.2, 3.05])
    add_picture(
        doc,
        hit_calib,
        6.45,
        "Figure 3. Retrieval rank sweep and confidence calibration from the saved evaluation run.",
    )
    add_para(
        doc,
        f"The separate retrieval-only rank sweep is more diagnostic than a single score: Hit@1 is {pct(hit[1]['rate'])} ({hit[1]['hits']}/{hit[1]['n']}), Hit@3 is {pct(hit[3]['rate'])} ({hit[3]['hits']}/{hit[3]['n']}), Hit@5 is {pct(hit[5]['rate'])} ({hit[5]['hits']}/{hit[5]['n']}), and Hit@10 is {pct(hit[10]['rate'])} ({hit[10]['hits']}/{hit[10]['n']}). The correct evidence is usually retrieved, but not always ranked first."
    )

    add_page_break(doc)
    add_heading(doc, "4. Final Results (continued)", 1)
    add_picture(
        doc,
        FIGURES / "fig4_score_hist.png",
        5.9,
        "Figure 4. Top-1 retrieval confidence split by Hit@1.",
    )
    add_picture(
        doc,
        heat_labels,
        6.45,
        "Figure 5. Hit@5 by answer label and the final label distribution.",
    )
    lora_rows = [
        ["Train / eval examples", f"{lora_train['train_examples']} / {lora_train['eval_examples']}", "Generated from indexed chunks"],
        ["Base model", lora_train["base_model"], "Small seq2seq model for local experiment"],
        ["Training loss metric", f"eval_loss = {fmt(lora_train['metrics']['eval_loss'])}", "After 3 epochs"],
        ["Adapter evaluation", f"mean Jaccard = {fmt(lora_eval['mean_jaccard_pred_vs_gold'])}", f"{lora_eval['examples']} generated eval rows"],
        ["QLoRA readiness", "false", "bitsandbytes=false, CUDA=false"],
    ]
    add_table(doc, ["LoRA/QLoRA result", "Value", "Meaning"], lora_rows, [1.65, 1.75, 2.9])
    add_para(
        doc,
        "The LoRA adapter result is useful as a style-adaptation experiment: it learns the context-answer-citation pattern in a small number of trainable weights. It does not replace the RAG pipeline, because the adapter cannot automatically know newly uploaded PDFs unless their evidence is still retrieved and supplied as context."
    )
    add_heading(doc, "LoRA Adapter Test", 2)
    add_para(
        doc,
        "After the midterm report, I added a LoRA / QLoRA experiment to compare the original FLAN-T5 model against the same FLAN-T5 model with my trained LoRA adapter attached. For the test, I gave both models the same question and the same retrieved course-document context, then compared whether the adapted model followed the grounded answer-and-citation style better than the base model."
    )
    run_rows = [
        [
            "Course project late-submission question",
            "Base FLAN-T5 answered only: No.",
            "The LoRA adapter answered with the retrieved syllabus text and citation, including: No late submission is permitted for course project-related deadlines.",
        ],
        [
            "Held-out syllabus evaluation question",
            "Base FLAN-T5 gave a short fragment: The assignment is adapted for the assignment.",
            "The LoRA adapter copied the relevant context much more completely and stayed closer to the expected answer.",
        ],
    ]
    add_table(doc, ["Question / case", "Base model answer", "LoRA adapter answer"], run_rows, [1.65, 2.15, 2.5])
    add_para(
        doc,
        "This did not mean the LoRA model became a full chatbot by itself. The important result was that fine-tuning helped the small model follow the context-and-citation style better than the base model. The main chatbot still needs RAG because new PDFs can be uploaded at any time."
    )

    add_page_break(doc)
    add_heading(doc, "5. Qualitative Examples and Comparative Interpretation", 1)
    example_rows = [
        [
            "Correct",
            "What is the late assignment policy?",
            "Late assignments are accepted up to three days after the due date with a penalty of 10% per day... [CSE_434___CSE_534_Syllabus.pdf p.2-7]",
            "top=0.695; sem=0.765; BERT=0.783",
        ],
        [
            "Correct with source conflict",
            "How long is Final Project Presentation time?",
            "The Final Project Presentation time is 8 minutes... while the Project Requirement document states 10-12 minutes...",
            "top=0.757; sem=0.824; BERT=0.779",
        ],
        [
            "Unsupported",
            "How many points is Assignment 1 worth?",
            "I could not find a supported answer in the uploaded documents.",
            "top=0.561; grounded=False",
        ],
        [
            "Partially Correct",
            "When is the midterm exam?",
            "The course does not include a midterm exam [CSE_434___CSE_534_Syllabus.pdf p.1].",
            "top=0.579; sem=0.580; BERT=0.773",
        ],
    ]
    add_table(
        doc,
        ["Label", "Question", "Actual saved output snippet", "Scores"],
        example_rows,
        [1.05, 1.55, 2.75, 0.95],
    )
    add_para(
        doc,
        "These cases show the main behavior pattern. The guardrail prevents unsupported invention, and citations make correct answers auditable. The remaining errors are mostly partially correct answers, where retrieval found relevant material but generation resolved ambiguous evidence poorly or paraphrased the gold answer too loosely."
    )
    add_para(
        doc,
        "The presentation deck frames the comparison qualitatively: with RAG, answers are cited and can refuse; without RAG, a raw LLM has no reliable course-document source and may invent specifics; with LoRA, the model can learn the answer style but remains limited unless retrieved context is still provided. No separate saved numeric raw-LLM baseline was present, so the report treats that comparison as qualitative rather than as an additional benchmark."
    )

    add_page_break(doc)
    add_heading(doc, "6. Short Reflection, Limitations & Future Work", 1)
    add_heading(doc, "Reflection", 2)
    add_para(
        doc,
        "The main lesson is that retrieval quality and refusal behavior mattered more than simply upgrading the generator. Strict prompting plus a pre-LLM confidence threshold produced a 0% hallucination label rate in the saved evaluation. The project also changed since the midterm by adding hybrid retrieval, BERTScore, calibration/ECE, live report figures, and the LoRA/QLoRA tab."
    )
    add_heading(doc, "Limitations", 2)
    for item in [
        "The evaluation set has only 25 questions because each row needs a gold answer and manual label.",
        "Most remaining errors are partially correct, not hallucinated; better generation checking or reranking is needed.",
        "pypdf extraction fails on scanned or unusual-layout PDFs unless OCR is added before ingestion.",
        "The default MiniLM embedding model is fast but can under-rank technical or long-context evidence.",
        "QLoRA could not be run locally because this machine lacks CUDA and bitsandbytes.",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "Future Work", 2)
    for item in [
        "Add OCR fallback for scanned documents.",
        "Add a cross-encoder reranker after FAISS/BM25 so the correct chunk is more often ranked first.",
        "Expand the evaluation set with more human-written questions and independent raters.",
        "Cache evaluation answers by question and retrieved chunk ids to reduce repeated API cost.",
        "Compare the LoRA adapter and RAG answers with a larger held-out set using semantic similarity and human judgments.",
    ]:
        add_bullet(doc, item)
    # Metadata for traceability.
    doc.core_properties.title = "Final Project Report - Generative AI RAG Document Q&A Chatbot"
    doc.core_properties.author = "Bhuwan Jung Thapa"
    doc.core_properties.subject = "CSE 434 / CSE 534 Generative AI final project report"
    doc.core_properties.keywords = "RAG, FAISS, BM25, LoRA, BERTScore, calibration, final project"

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    make_report()
